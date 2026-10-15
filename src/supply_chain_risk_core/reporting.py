"""从同一分析结果生成 Markdown、网页、矢量图和 Office 结果快照。"""
from dataclasses import dataclass
from datetime import date
from html import escape
from io import BytesIO
import math
from pathlib import Path
import re
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from .credit.labels import STATUS_LABELS, input_label, rule_label
from .validation import canonical_bytes


@dataclass(frozen=True)
class Section:
    title: str
    paragraphs: tuple = ()
    headers: tuple = ()
    rows: tuple = ()


def display(value):
    if value is None:
        return "未量化"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float):
        return f"{value:,.4f}".rstrip("0").rstrip(".")
    return str(value)


def _scenario_rows(chain):
    result = chain["scenario"]["result"]
    common = (("冲击前 EBIT", "ebit_before_yi", "亿元"), ("冲击后 EBIT", "ebit_after_yi", "亿元"),
              ("冲击前利息保障倍数", "interest_cover_before_x", "倍"),
              ("冲击后利息保障倍数", "interest_cover_after_x", "倍"), ("临界冲击幅度", "critical_shock_pct", "%"))
    amount = (("新增成本", "added_cost_yi", "亿元"),) if chain["scenario"]["kind"] == "cost_push" else (
        ("收入下降", "revenue_loss_yi", "亿元"), ("EBIT 损失", "ebit_loss_yi", "亿元"))
    return tuple((label, result[key], unit) for label, key, unit in amount + common)


def sections(report):
    credit = report["credit"]["companies"]
    rows = tuple((company["company_id"], company["name"], company["action"], company["risk_score"],
                  "需复核" if company["review_required"] else "输入已齐", "；".join(company["risk_control_advice"]) or "无新增措施")
                 for company in credit)
    result = [Section("授信模型结果", ("风险分采用原参赛规则的绝对刻度。缺失输入保留在明细中，不能据此认定没有风险。",),
                      ("编号", "企业", "模型方向", "风险分", "输入状态", "风控建议"), rows)]
    for company in credit:
        basis = company["hard_rule_hits"] + company["combination_hits"] + company["financial_rule_hits"]
        result.append(Section(f"{company['company_id']} 决策依据",
                              (f"触发规则：{'；'.join(map(rule_label, basis)) or '未命中收紧规则'}。",
                               f"未通过的放宽条件：{'；'.join(map(rule_label, company['loosen_gate_failures'])) or '无'}。",
                               f"未知输入：{'、'.join(map(input_label, company['missing_inputs'])) or '无'}。",
                               f"建议证据状态：{STATUS_LABELS[company['model_advice_assessment']['status']]}。"),
                              ("主体资质", "风险信号", "行业", "区域"),
                              (tuple(company["scores"][key] for key in ("score_entity", "score_signal", "score_industry", "score_region")),)))
    for chain in report["chains"]:
        parameters = chain["scenario"]["parameters"]
        if chain["scenario"]["kind"] == "cost_push":
            assumptions = f"情景参数：源头价格上涨{parameters['shock']*100:g}%，转嫁率{parameters['pass_through']*100:g}%，源头价格影响系数{parameters['coefficient']:g}。"
        else:
            margin = "未知" if parameters["contribution_margin"] is None else f"{parameters['contribution_margin']*100:g}%"
            assumptions = f"情景参数：需求下降{parameters['shock']*100:g}%，销售份额{parameters['sales_share']*100:g}%，贡献率{margin}。"
        result.append(Section(chain["title"],
                              (chain["relation_text"], chain["path_text"], assumptions, chain["impact_text"], chain["scenario"]["result"]["verdict"]),
                              ("情景指标", "结果", "单位"), _scenario_rows(chain)))
        result.append(Section(f"{chain['title']}的节点比较", (chain["node_text"],),
                              ("节点", "影响范围", "不可替代性", "传导后果", "风险值"),
                              tuple((row["name"], row["scores"]["impact"], row["scores"]["substitutability"],
                                     row["scores"]["consequence"], row["risk_value"]) for row in chain["comparison_rows"])))
    labels = {"retained_cost_per_100_revenue": "每百元基期收入吸收成本", "cost_change_pct": "采购组合成本变化",
              "board_transfer_fraction": "板卡转嫁率", "brand_transfer_fraction": "品牌转嫁率", "board_retained_cost": "板卡吸收额",
              "brand_retained_cost": "品牌吸收额", "downstream_cost": "下游承接额", "brand_minus_board_cost": "品牌与板卡吸收差额",
              "comparison": "比较结论", "lower": "交叉收入下界", "upper": "交叉收入上界"}
    for case in report["quantitative_cases"]:
        values = []
        for key, value in case["result"].items():
            if key == "unit":
                continue
            unit = case["result"]["unit"]
            if key.endswith("_fraction"):
                value, unit = value * 100, "%"
            elif isinstance(value, str):
                unit = "—"
            values.append((labels[key], value, unit))
        result.append(Section(case["title"], (), ("指标", "结果", "单位"), tuple(values)))
    io = report["input_output"]
    result.append(Section("投入产出计算", (io["coefficient_basis"], io["price_basis"]),
                          ("行业", "源头价格变化系数"), tuple(io["source_price_coefficients"].items())))
    result.append(Section("来源与适用边界", (report["source_verification"], report["interpretation"]),
                          ("来源", "公开日期", "链接"),
                          tuple((source["title"], source["published_at"], source["url"]) for source in report["sources"])))
    return result


def _intro(report):
    kind = "演示数据，企业与数值为合成输入" if report["data_kind"] == "synthetic_demonstration" else "公开证据分析"
    return f"{kind}。评估基准日：{report['as_of']}。"


def _markdown(value):
    return display(value).replace("|", "\\|").replace("\n", "<br>")


def markdown_report(report):
    lines = [f"# {report['title']}", "", _intro(report), ""]
    for section in sections(report):
        lines.extend([f"## {section.title}", "", *section.paragraphs, ""])
        if section.headers:
            lines.append("| " + " | ".join(section.headers) + " |")
            lines.append("| " + " | ".join("---" for _ in section.headers) + " |")
            lines.extend("| " + " | ".join(_markdown(value) for value in row) + " |" for row in section.rows)
            lines.append("")
    return "\n".join(lines)


def _svg_lines(value, width):
    lines, line, units = [], "", 0
    for char in value:
        size = 2 if ord(char) > 255 else 1
        if units + size > width and line:
            lines.append(line)
            line, units = "", 0
        line += char
        units += size
    return lines + ([line] if line else [])


def chain_svg(chain):
    layers = chain["layers"]
    width = 1100
    label_lines = {node["id"]: _svg_lines(node["label"].replace("\n", ""), 14)
                   for layer in layers for node in layer}
    box_height = max(64, max(len(lines) for lines in label_lines.values()) * 20 + 22)
    layer_count = max(len(layer) for layer in layers)
    pitch = box_height + 28
    graph_bottom = 90 + layer_count * pitch
    caption = _svg_lines(chain["mechanism_label"] + "：" + chain["transmission_label"], 122)
    details = "；".join(f"{step['from_name']}→{step['to_name']} {step['level_detail']}" for step in chain["path_steps"])
    detail_lines = _svg_lines(details, 138)
    height = graph_bottom + 84 + len(caption) * 24 + len(detail_lines) * 22
    positions = {node["id"]: (90 + index * 900 / max(1, len(layers) - 1),
                              90 + (layer_count - len(layer)) * pitch / 2 + offset * pitch + box_height / 2)
                 for index, layer in enumerate(layers) for offset, node in enumerate(layer)}
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}" role="img">',
             f'<title>{escape(chain["title"])}</title>', f'<desc>{escape(chain["transmission_label"])}</desc>',
             '<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3z" fill="#466481"/></marker></defs>',
             f'<rect width="1100" height="{height}" rx="16" fill="#f4f7fb"/>',
             '<g font-family="Microsoft YaHei,Noto Sans CJK SC,sans-serif">',
             f'<text x="34" y="44" font-size="23" fill="#172b45">{escape(chain["title"])}</text>']
    for source, target in chain["links"]:
        x1, y1 = positions[source]
        x2, y2 = positions[target]
        parts.append(f'<path d="M{x1+65},{y1} L{x2-72},{y2}" fill="none" stroke="#9aacc0" stroke-width="2" marker-end="url(#arrow)"/>')
    for layer in layers:
        for node in layer:
            x, y = positions[node["id"]]
            fill = "#fbe7e6" if node["id"] == chain["highest_node"] else "#e2f3ec" if node["id"] == chain["lowest_node"] else "#ffffff"
            parts.append(f'<rect x="{x-66}" y="{y-box_height/2}" width="132" height="{box_height}" rx="10" fill="{fill}" stroke="#a2b4c6"/>')
            lines = label_lines[node["id"]]
            for index, line in enumerate(lines):
                y_text = y - (len(lines) - 1) * 10 + index * 20 + 5
                parts.append(f'<text data-node="{node["id"]}" x="{x}" y="{y_text}" text-anchor="middle" font-size="16" fill="#172b45">{escape(line)}</text>')
    for index, line in enumerate(caption):
        parts.append(f'<text x="34" y="{graph_bottom+28+index*24}" font-size="16" fill="#304b66">{escape(line)}</text>')
    for index, line in enumerate(detail_lines):
        parts.append(f'<text x="34" y="{graph_bottom+42+len(caption)*24+index*22}" font-size="14" fill="#304b66">{escape(line)}</text>')
    parts.append(f'<text x="34" y="{height-18}" font-size="13" fill="#52687d">风险值仅用于链内比较；红色为最高节点，绿色为最低节点。</text></g></svg>')
    return "".join(parts)


def html_report(report):
    content = [f'<p class="intro">{escape(_intro(report))}</p>']
    diagrams = {chain["title"]: chain for chain in report["chains"]}
    for section in sections(report):
        content.append(f'<section><h2>{escape(section.title)}</h2>')
        if section.title in diagrams:
            content.append(chain_svg(diagrams[section.title]))
        content.extend(f'<p>{escape(paragraph).replace(chr(10), "<br>")}</p>' for paragraph in section.paragraphs)
        if section.headers:
            header = "".join(f"<th>{escape(cell)}</th>" for cell in section.headers)
            rows = "".join("<tr>" + "".join(f"<td>{escape(display(value))}</td>" for value in row) + "</tr>" for row in section.rows)
            content.append(f'<div class="table-scroll"><table><thead><tr>{header}</tr></thead><tbody>{rows}</tbody></table></div>')
        content.append("</section>")
    style = "body{margin:0;background:#edf2f8;color:#20334a;font:16px/1.75 'Microsoft YaHei',sans-serif}main{max-width:1100px;margin:40px auto;padding:42px;background:white;border-radius:18px}h1{font-size:32px;line-height:1.35;color:#152c46}h2{font-size:23px;margin:0 0 20px;color:#152c46}section{padding:28px 0;border-top:1px solid #dce5ef}p{overflow-wrap:anywhere}.intro{color:#526b84;margin-bottom:30px}svg{display:block;width:100%;height:auto;margin:20px 0}.table-scroll{overflow:auto}table{width:100%;border-collapse:collapse;font-size:14px}th{background:#1e3b5d;color:white;text-align:left}th,td{padding:13px 14px;border:1px solid #dce5ef;vertical-align:top}tr:nth-child(even){background:#f5f8fc}@media(max-width:700px){main{margin:0;padding:22px;border-radius:0}h1{font-size:27px}th,td{min-width:72px}}"
    return f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{escape(report["title"])}</title><style>{style}</style><main><h1>{escape(report["title"])}</h1>{"".join(content)}</main></html>'


def _normalize_office(path, as_of):
    stamp = f"{as_of}T00:00:00Z"
    buffer = BytesIO()
    with ZipFile(path) as source, ZipFile(buffer, "w", ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == "docProps/core.xml":
                xml = data.decode("utf-8")
                xml = re.sub(r"(<dcterms:(?:created|modified)[^>]*>).*?(</dcterms:(?:created|modified)>)", rf"\g<1>{stamp}\g<2>", xml)
                data = xml.encode("utf-8")
            member = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            member.compress_type = ZIP_DEFLATED
            target.writestr(member, data)
    Path(path).write_bytes(buffer.getvalue())


def export_docx(report, path):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

    doc = Document()
    doc.core_properties.author = "Hokkaidoooo"
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    for border in list(doc.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    doc.add_paragraph(report["title"], "Title")
    doc.add_paragraph(_intro(report))
    for section in sections(report):
        doc.add_heading(section.title, 1)
        for index, paragraph in enumerate(section.paragraphs):
            paragraph_node = doc.add_paragraph(paragraph)
            paragraph_node.paragraph_format.keep_with_next = index == len(section.paragraphs) - 1
        if not section.headers:
            continue
        table = doc.add_table(rows=1, cols=len(section.headers))
        table.style = "Table Grid"
        table.autofit = False
        if len(section.headers) == 6:
            widths = (1.7, 3.4, 1.7, 1.5, 1.8, 6.9)
        elif len(section.headers) == 3 and section.headers[0] == "来源":
            widths = (3.4, 3.0, 10.6)
        elif len(section.headers) == 3:
            widths = (7.4, 4.2, 5.4)
        elif len(section.headers) == 5:
            widths = (4.2, 3.2, 3.2, 3.2, 3.2)
        else:
            widths = tuple(17 / len(section.headers) for _ in section.headers)
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            edge = OxmlElement(f"w:{side}")
            for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                edge.set(qn(f"w:{attr}"), value)
            borders.append(edge)
        table._tbl.tblPr.append(borders)
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for index, heading in enumerate(section.headers):
            table.rows[0].cells[index].text = heading
        for values in section.rows:
            cells = table.add_row().cells
            for index, value in enumerate(values):
                cells[index].text = display(value)
        for index, row in enumerate(table.rows):
            cant_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant_split)
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "DCE6F1" if index == 0 else "F4F7FB" if index % 2 == 0 else "FFFFFF")
                properties.append(shading)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    edge = OxmlElement(f"w:{side}")
                    edge.set(qn("w:w"), "100")
                    edge.set(qn("w:type"), "dxa")
                    margins.append(edge)
                properties.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = index == 0
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)
                        run.bold = index == 0
    doc.save(path)
    _normalize_office(path, report["as_of"])


def export_xlsx(report, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    book.properties.creator = "Hokkaidoooo"
    groups = [("授信概览", sections(report)[:1]),
              ("决策依据", [section for section in sections(report) if section.title.endswith("决策依据")]),
              ("传导情景", [section for section in sections(report) if section.title in {chain["title"] for chain in report["chains"]}]),
              ("节点比较", [section for section in sections(report) if section.title.endswith("节点比较")]),
              ("定量与来源", [section for section in sections(report) if section.title in {case["title"] for case in report["quantitative_cases"]} or section.title in {"投入产出计算", "来源与适用边界"}])]
    edge = Side(style="thin", color="D9D9D9")
    widths_by_sheet = {"授信概览": (15, 26, 14, 12, 16, 58), "决策依据": (22, 22, 22, 22),
                       "传导情景": (38, 28, 20), "节点比较": (24, 18, 20, 18, 16), "定量与来源": (38, 28, 52)}
    for name, selected in groups:
        sheet = book.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.append([report["title"]])
        sheet.append([_intro(report)])
        sheet.append(["结果快照；修改输入后请重新生成报告。"])
        sheet.append([])
        widths = widths_by_sheet[name]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for section in selected:
            sheet.append([section.title])
            for paragraph in section.paragraphs:
                row = sheet.max_row + 1
                sheet.append([paragraph])
                sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(widths))
                units = sum(widths) - 4
                count = sum(max(1, math.ceil(sum(2 if ord(char) > 255 else 1 for char in line) / units)) for line in paragraph.split("\n"))
                sheet.row_dimensions[row].height = 18 * count + 10
            if section.headers:
                sheet.append(list(section.headers))
                for cell in sheet[sheet.max_row][:len(section.headers)]:
                    cell.fill = PatternFill("solid", fgColor="1E3B5D")
                    cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True, size=10)
                for values in section.rows:
                    converted = [date.fromisoformat(value) if section.headers[index] == "公开日期" else "未量化" if value is None else value
                                 for index, value in enumerate(values)]
                    sheet.append(converted)
            sheet.append([])
        merged_rows = {area.min_row for area in sheet.merged_cells.ranges}
        for row in sheet:
            for cell in row:
                if cell.value is None:
                    continue
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                if cell.fill.fgColor.rgb != "001E3B5D":
                    cell.font = Font(name="Microsoft YaHei", size=10, color="20334A")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0' if float(cell.value).is_integer() else '#,##0.00##'
                elif isinstance(cell.value, date):
                    cell.number_format = 'yyyy-mm-dd'
            row_number = row[0].row
            if row_number not in merged_rows:
                lines = []
                for cell in row:
                    value = display(cell.value) if cell.value is not None else ""
                    width = sheet.column_dimensions[cell.column_letter].width - 2
                    lines.append(sum(max(1, math.ceil(sum(2 if ord(char) > 255 else 1 for char in line) / width))
                                     for line in value.split("\n")))
                sheet.row_dimensions[row_number].height = max(34, 17 * max(lines) + 10)
        sheet.row_dimensions[1].height = 34
        sheet.row_dimensions[2].height = 28
        sheet.row_dimensions[3].height = 24
        sheet["A1"].font = Font(name="Microsoft YaHei", size=17, bold=True, color="172B45")
        for row in (1, 2, 3):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(widths))
        sheet.freeze_panes = "B5"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.print_options.horizontalCentered = True
        sheet.print_title_rows = "1:3"
    book.save(path)
    _normalize_office(path, report["as_of"])


def write_reports(report, directory, *, office=True):
    root = Path(directory)
    (root / "analysis.json").write_bytes(canonical_bytes(report) + b"\n")
    (root / "report.md").write_text(markdown_report(report), encoding="utf-8", newline="\n")
    (root / "report.html").write_text(html_report(report), encoding="utf-8", newline="\n")
    for index, chain in enumerate(report["chains"], 1):
        (root / f"chain-{index:02d}.svg").write_text(chain_svg(chain), encoding="utf-8", newline="\n")
    if office:
        export_docx(report, root / "report.docx")
        export_xlsx(report, root / "report.xlsx")
